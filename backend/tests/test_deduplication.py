import unittest
import uuid
from datetime import UTC, datetime
from io import BytesIO
from unittest.mock import patch

from docx import Document
from fastapi.testclient import TestClient

from app.config import get_settings
from app.main import app
from app.services.deduplication_service import DeduplicationService
from app.services.hash_service import calculate_sha256
from app.services.similarity_service import find_most_similar
from app.services.text_extraction import extract_text


class MemoryResult:
    def __init__(self, data):
        self.data = data


class MemoryQuery:
    def __init__(self, client, table):
        self.client = client
        self.name = table
        self.action = "select"
        self.filters = []
        self.ordering = None
        self.limit_value = None
        self.range_value = None
        self.payload = None
        self.return_single = False

    def select(self, *_args, **_kwargs): return self
    def eq(self, key, value): self.filters.append((key, value)); return self
    def neq(self, key, value): self.filters.append((key, ("neq", value))); return self
    def in_(self, key, values): self.filters.append((key, ("in", values))); return self
    def order(self, key, desc=False): self.ordering = (key, desc); return self
    def limit(self, value): self.limit_value = value; return self
    def range(self, start, end): self.range_value = (start, end); return self
    def maybe_single(self): self.return_single = True; return self
    def single(self): self.return_single = True; return self
    def insert(self, payload): self.action = "insert"; self.payload = payload; return self
    def update(self, payload): self.action = "update"; self.payload = payload; return self
    def upsert(self, payload, **_kwargs): self.action = "upsert"; self.payload = payload; return self
    def delete(self): self.action = "delete"; return self

    def execute(self):
        rows = self.client.tables.setdefault(self.name, [])
        if self.action == "insert":
            values = self.payload if isinstance(self.payload, list) else [self.payload]
            inserted = []
            for value in values:
                row = {"id": str(uuid.uuid4()), "created_at": datetime.now(UTC).isoformat(), **value}
                rows.append(row)
                inserted.append(row.copy())
            return MemoryResult(inserted)
        matches = []
        for row in rows:
            if all(
                (row.get(key) != value[1]) if isinstance(value, tuple) and value[0] == "neq"
                else (row.get(key) in value[1]) if isinstance(value, tuple) and value[0] == "in"
                else row.get(key) == value
                for key, value in self.filters
            ):
                matches.append(row)
        if self.ordering:
            key, desc = self.ordering
            matches.sort(key=lambda row: row.get(key) or "", reverse=desc)
        if self.range_value:
            matches = matches[self.range_value[0]:self.range_value[1] + 1]
        if self.limit_value is not None:
            matches = matches[:self.limit_value]
        if self.action == "update":
            for row in matches:
                row.update(self.payload)
            return MemoryResult([row.copy() for row in matches])
        if self.action == "upsert":
            existing = next((row for row in rows if row.get("file_id") == self.payload.get("file_id")), None)
            if existing:
                existing.update(self.payload)
                return MemoryResult([existing.copy()])
            rows.append(self.payload.copy())
            return MemoryResult([self.payload.copy()])
        if self.action == "delete":
            self.client.tables[self.name] = [row for row in rows if row not in matches]
            return MemoryResult([])
        if self.return_single:
            return MemoryResult(matches[0].copy() if matches else None)
        return MemoryResult([row.copy() for row in matches])


class MemoryBucket:
    def __init__(self, storage): self.storage = storage
    def upload(self, path, content, **_kwargs):
        if isinstance(content, str):
            with open(content, "rb") as source:
                content = source.read()
        self.storage.objects[path] = content.read() if hasattr(content, "read") else content
    def remove(self, paths):
        for path in paths: self.storage.objects.pop(path, None)
    def download(self, path): return self.storage.objects[path]


class MemoryStorage:
    def __init__(self): self.objects = {}
    def from_(self, _bucket): return MemoryBucket(self)


class MemorySupabase:
    def __init__(self):
        self.tables = {"files": [], "file_analysis": []}
        self.storage = MemoryStorage()
    def table(self, name): return MemoryQuery(self, name)


def make_pdf(text: str) -> BytesIO:
    stream_content = f"BT /F1 12 Tf 72 720 Td ({text}) Tj ET".encode("ascii")
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Resources << /Font << /F1 4 0 R >> >> /Contents 5 0 R >>",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
        b"<< /Length " + str(len(stream_content)).encode() + b" >>\nstream\n" + stream_content + b"\nendstream",
    ]
    pdf = bytearray(b"%PDF-1.4\n")
    offsets = [0]
    for index, obj in enumerate(objects, 1):
        offsets.append(len(pdf))
        pdf.extend(f"{index} 0 obj\n".encode() + obj + b"\nendobj\n")
    xref = len(pdf)
    pdf.extend(f"xref\n0 {len(objects) + 1}\n0000000000 65535 f \n".encode())
    for offset in offsets[1:]:
        pdf.extend(f"{offset:010d} 00000 n \n".encode())
    pdf.extend(f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF".encode())
    return BytesIO(bytes(pdf))


class IntelligenceTests(unittest.TestCase):
    def setUp(self):
        self.client = MemorySupabase()
        self.settings = get_settings()

    def test_hash_is_binary_based_and_restores_pointer(self):
        stream = BytesIO(b"same bytes")
        stream.seek(4)
        first = calculate_sha256(stream)
        self.assertEqual(stream.tell(), 4)
        self.assertEqual(first, calculate_sha256(BytesIO(b"same bytes")))
        self.assertNotEqual(first, calculate_sha256(BytesIO(b"different bytes")))

    def test_exact_duplicate_avoids_second_storage_object(self):
        service = DeduplicationService(self.client, self.settings)
        original = b"Cloud storage keeps project documents safely."
        first = service.analyze(BytesIO(original), "first.txt")
        self.assertEqual(first.deduplication_status, "unique")
        self.client.tables["files"].append({
            "id": "original-id", "file_hash": first.file_hash, "duplicate_of": None,
            "analysis_status": "completed", "created_at": "2026-01-01",
        })
        second = service.analyze(BytesIO(original), "copy.txt")
        self.assertEqual(second.deduplication_status, "duplicate")
        self.assertEqual(second.duplicate_of, "original-id")
        self.assertEqual(second.similarity_score, 1.0)
        self.assertTrue(second.storage_saved)
        self.client.tables["files"].append({
            "id": "duplicate-id", "file_hash": first.file_hash, "duplicate_of": "original-id",
            "analysis_status": "completed", "created_at": "2026-01-02",
        })
        rerun = service.analyze(BytesIO(original), "first.txt", exclude_file_id="original-id")
        self.assertEqual(rerun.deduplication_status, "unique")

    def test_similarity_and_unrelated_classification(self):
        left = "cloud computing provides scalable storage for reliable project data secure document access across modern business networks"
        right = "cloud computing provides scalable storage for reliable project data secure file access across modern business networks"
        match = find_most_similar(left, [("candidate", right)])
        self.assertIsNotNone(match)
        self.assertGreaterEqual(match.score, self.settings.similarity_threshold)
        self.assertNotEqual(calculate_sha256(BytesIO(left.encode())), calculate_sha256(BytesIO(right.encode())))

        unrelated = find_most_similar(
            "Cloud computing provides scalable storage for reliable project data.",
            [("candidate", "Volcanic geology examines magma, tectonic eruptions, and mineral formations.")],
        )
        self.assertTrue(unrelated is None or unrelated.score < self.settings.similarity_threshold)

        self.client.tables["files"].append({
            "id": "pdf-candidate", "file_hash": "unrelated-hash", "duplicate_of": None,
            "analysis_status": "completed", "created_at": "2026-01-01",
        })
        self.client.tables["file_analysis"].append({
            "file_id": "pdf-candidate", "extracted_text": left,
        })
        docx_result = DeduplicationService(self.client, self.settings).analyze(
            BytesIO(right.encode()), "similar.txt"
        )
        self.assertEqual(docx_result.deduplication_status, "similar")
        self.assertEqual(docx_result.duplicate_of, "pdf-candidate")

    def test_low_similarity_keeps_numeric_score_on_unique_result(self):
        candidate_text = "Cloud computing supports resilient storage and document collaboration for distributed teams."
        self.client.tables["files"].append({
            "id": "candidate-id", "file_hash": "candidate-hash", "duplicate_of": None,
            "analysis_status": "completed", "created_at": "2026-01-01",
        })
        self.client.tables["file_analysis"].append({
            "file_id": "candidate-id", "extracted_text": candidate_text,
        })
        new_text = "Weather forecasting studies ocean currents, cloud formations, and atmospheric pressure patterns."
        result = DeduplicationService(self.client, self.settings).analyze(
            BytesIO(new_text.encode()), "new.txt"
        )
        self.assertEqual(result.deduplication_status, "unique")
        self.assertIsNotNone(result.similarity_score)
        self.assertLess(result.similarity_score, self.settings.similarity_threshold)
        self.assertIsNone(result.duplicate_of)

    def test_pdf_docx_and_unsupported_extraction(self):
        shared_text = "Cloud computing provides scalable storage for reliable project data secure document access across modern business networks"
        pdf_stream = make_pdf(shared_text)
        pdf_text, pdf_status = extract_text(pdf_stream, "sample.pdf", 30000)
        self.assertEqual(pdf_status, "completed")
        self.assertIn("cloud computing", pdf_text)

        doc = Document()
        doc.add_paragraph(shared_text.replace("document", "file"))
        docx_stream = BytesIO()
        doc.save(docx_stream)
        docx_text, docx_status = extract_text(docx_stream, "sample.docx", 30000)
        self.assertEqual(docx_status, "completed")
        self.assertIn("cloud computing", docx_text)
        self.client.tables["files"].append({
            "id": "pdf-candidate", "file_hash": "different-pdf-hash", "duplicate_of": None,
            "analysis_status": "completed", "created_at": "2026-01-01",
        })
        self.client.tables["file_analysis"].append({
            "file_id": "pdf-candidate", "extracted_text": pdf_text,
        })
        docx_result = DeduplicationService(self.client, self.settings).analyze(
            docx_stream, "sample.docx"
        )
        self.assertEqual(docx_result.deduplication_status, "similar")
        self.assertEqual(docx_result.duplicate_of, "pdf-candidate")

        unsupported_text, unsupported_status = extract_text(BytesIO(b"\x00\x01binary"), "image.png", 30000)
        self.assertEqual((unsupported_text, unsupported_status), ("", "unsupported"))

    def test_api_upload_duplicate_and_step_one_routes(self):
        from app.api.routes import files as file_routes

        memory = self.client
        with patch.object(file_routes, "get_supabase_client", return_value=memory):
            api = TestClient(app)
            health = api.get("/api/health")
            self.assertEqual(health.status_code, 200)
            first = api.post("/api/files/upload", files={"file": ("one.txt", b"Exact content for storage.", "text/plain")})
            self.assertEqual(first.status_code, 201, first.text)
            duplicate = api.post("/api/files/upload", files={"file": ("copy.txt", b"Exact content for storage.", "text/plain")})
            self.assertEqual(duplicate.status_code, 201, duplicate.text)
            self.assertEqual(duplicate.json()["file"]["deduplication_status"], "duplicate")
            self.assertEqual(len(memory.storage.objects), 1)
            listing = api.get("/api/files")
            self.assertEqual(listing.status_code, 200, listing.text)
            self.assertEqual(len(listing.json()["files"]), 2)
            detail = api.get(f"/api/files/{first.json()['file']['id']}")
            self.assertEqual(detail.status_code, 200, detail.text)
            reanalysis = api.post(f"/api/files/{first.json()['file']['id']}/analyze")
            self.assertEqual(reanalysis.status_code, 200, reanalysis.text)
            self.assertEqual(reanalysis.json()["file"]["deduplication_status"], "unique")
            self.assertEqual(api.get("/api/files/not-a-uuid").status_code, 400)
            self.assertEqual(api.post("/api/files/upload", data={}).status_code, 400)
            self.assertEqual(api.post("/api/files/upload", files={"file": ("empty.txt", b"", "text/plain")}).status_code, 400)
            small_limit = self.settings.model_copy(update={"max_upload_size_bytes": 4})
            with patch.object(file_routes, "get_settings", return_value=small_limit):
                oversized = api.post("/api/files/upload", files={"file": ("big.txt", b"12345", "text/plain")})
            self.assertEqual(oversized.status_code, 413)
            self.assertEqual(api.get("/api/files/duplicates").status_code, 200)
            analytics = api.get("/api/analytics")
            self.assertEqual(analytics.status_code, 200, analytics.text)
            self.assertEqual(analytics.json()["estimated_storage_saved_bytes"], len(b"Exact content for storage."))


if __name__ == "__main__":
    unittest.main()
